"""Generate synthetic Copilot-style transcript .docx files.

Everything here is fictional (a WideWorldImporters-flavored stand-in). The layout follows
parse_docx's ASSUMED Copilot format; regenerate once the kb-survey report confirms the real one.

Usage: python fixtures/make_fixtures.py OUT_DIR
"""

import sys
from pathlib import Path

from docx import Document

TRANSCRIPTS = {
    "chatbot-requirements-sync.docx": (
        ["Chatbot requirements sync", "September 8, 2026, 10:00AM", "12m 30s"],
        [
            ("Sam Rivera", "0:05", ["Okay, um, so I think everyone's here. John, you wanted to go over what the chatbot should actually do for v1?"]),
            ("John Doe", "0:15", ["Yeah. So, uh, the main thing is customers need to ask about order status. Like, where's my order, has it shipped, that kind of thing."]),
            ("John Doe", "0:32", ["Second, um, invoice history. They should be able to say show me my last five invoices and, uh, get that back."]),
            ("John Doe", "0:48", ["And third, if the bot can't answer, it has to hand off to a human agent. That's, that's non-negotiable."]),
            ("Marcus Lee", "1:05", ["Got it. What about integrations? Where does this thing live?"]),
            ("John Doe", "1:12", ["It lives in Teams first. Um, and it needs to open tickets in ServiceNow when it hands off to a human."]),
            ("Marcus Lee", "1:30", ["Do we need Salesforce? Sales has been asking about pulling account data."]),
            ("John Doe", "1:41", ["No, uh, not for v1. Salesforce is a phase two thing. Let's not scope creep."]),
            ("Sam Rivera", "1:55", ["Okay. And the invoice history, that's coming from the Sales schema, right? Sales.Invoices?"]),
            ("Marcus Lee", "2:06", ["Yeah, Sales.Invoices and InvoiceLines. I'll, uh, I'll write a read-only proc for it."]),
            ("John Doe", "2:20", ["Great. Can we get a demo by end of month?"]),
            ("Marcus Lee", "2:25", ["End of month is tight but, um, yeah, I'll have a demo by September 30th."]),
        ],
    ),
    "kt-customer-bi.docx": (
        ["KT: Customer BI pipeline", "September 15, 2026, 2:00PM", "9m 05s"],
        [
            ("Sam Rivera", "0:03", ["Thanks for doing this, Priya. Um, so the big question: where does customer BI actually get calculated?"]),
            ("Priya Raman", "0:12", ["So, uh, it's all in Reporting.usp_CalcCustomerBI. That proc reads Sales.Invoices, Sales.InvoiceLines and Sales.Customers, and it writes into Reporting.CustomerBISnapshot."]),
            ("Sam Rivera", "0:35", ["And what kicks it off?"]),
            ("Priya Raman", "0:39", ["Reporting.usp_RefreshBIDaily calls it. That runs, um, nightly from a SQL Agent job at 2 AM."]),
            ("Sam Rivera", "0:52", ["Why is it a proc? I thought BI used to be a view."]),
            ("Priya Raman", "1:01", [
                "It was. Back in 2024 the calc lived in Reporting.vw_CustomerBI, but, uh, the view kept timing out once invoice volume grew.",
                "So we moved the calc into the proc and the view just reads the snapshot now.",
            ]),
            ("Sam Rivera", "1:25", ["Got it. Anything that depends on the snapshot that I should know about?"]),
            ("Priya Raman", "1:31", ["Yeah, the big one is Integration.usp_ExportBIFeed. It builds its query with dynamic SQL against the view, so, um, the dependency tools don't see it."]),
            ("Sam Rivera", "1:48", ["Who owns the export feed?"]),
            ("Priya Raman", "1:53", ["Honestly, I don't know. It was Dave's, and he left. Someone needs to figure that out."]),
            ("Sam Rivera", "2:05", ["Okay, I'll flag that."]),
        ],
    ),
    "chatbot-sprint-review.docx": (
        ["Chatbot sprint review", "September 22, 2026, 11:00AM", "8m 40s"],
        [
            ("Sam Rivera", "0:04", ["Alright, uh, sprint review. Marcus, how's the demo looking?"]),
            ("Marcus Lee", "0:12", ["Order status works end to end. Invoice history is, um, mostly there."]),
            ("John Doe", "0:25", ["Nice. So, uh, change of plans on one thing. Leadership wants Salesforce in v1 after all."]),
            ("John Doe", "0:38", ["They want the bot to pull the customer's account manager from Salesforce, so, um, the handoff goes to the right person."]),
            ("Marcus Lee", "0:55", ["That's a real scope change. I'd need to look at the Salesforce API first."]),
            ("John Doe", "1:04", ["Understood. Can you scope it by October 3rd?"]),
            ("Marcus Lee", "1:10", ["Yeah, I'll have a Salesforce scoping doc by October 3rd."]),
            ("Sam Rivera", "1:20", ["Does that push the demo?"]),
            ("John Doe", "1:26", ["Keep the demo on the 30th without Salesforce. Uh, Salesforce comes after."]),
            ("Sam Rivera", "1:38", ["And ServiceNow tickets still happen on handoff?"]),
            ("John Doe", "1:44", ["Yes, nothing changes there."]),
        ],
    ),
    "data-platform-standup.docx": (
        ["Data platform standup", "October 6, 2026, 9:30AM", "4m 10s"],
        [
            ("Lena Ortiz", "0:02", ["Quick one from me. The BI snapshot was late again on Saturday."]),
            ("Sam Rivera", "0:09", ["Late how?"]),
            ("Lena Ortiz", "0:12", ["The invoice load ran long, so, um, usp_RefreshBIDaily started at 2 like always but CustomerBISnapshot didn't finish until almost 6."]),
            ("Marcus Lee", "0:30", ["Does the chatbot read that? Because invoice history in the bot hits Sales.Invoices directly, not the snapshot."]),
            ("Lena Ortiz", "0:42", ["Right, the bot's fine. It's the export feed that got stale data."]),
            ("Sam Rivera", "0:50", ["Which still has no owner, by the way."]),
            ("Lena Ortiz", "0:56", ["I can take ownership of the export feed. I'll, uh, write up how it works by October 17th."]),
            ("Sam Rivera", "1:08", ["Perfect, thank you."]),
        ],
    ),
}


def write(out_dir: Path) -> list[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    written = []
    for name, (preamble, turns) in TRANSCRIPTS.items():
        doc = Document()
        for line in preamble:
            doc.add_paragraph(line)
        for speaker, ts, paragraphs in turns:
            doc.add_paragraph(f"{speaker}   {ts}")
            for text in paragraphs:
                doc.add_paragraph(text)
        path = out_dir / name
        doc.save(str(path))
        written.append(path)
    return written


if __name__ == "__main__":
    for p in write(Path(sys.argv[1])):
        print(p)
